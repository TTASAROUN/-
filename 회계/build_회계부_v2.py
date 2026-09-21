# -*- coding: utf-8 -*-
"""따사로운 회계부 v2 — 사업자별(통합서비스 / 평생교육원) 완전 분리 구조."""
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L

OUT = "따사로운_회계부_v2_사업자별.xlsx"
FONT = "맑은 고딕"
NROW0, NROW1 = 15, 2014          # 거래입력 데이터 영역
NUM = '#,##0;[Red]-#,##0;"-"'

# ---------- 색 ----------
C_TXT, C_SUB, C_MUTE = "414042", "58595B", "A7A9AC"
C_HEAD = "37383A"
C_RED, C_GREEN, C_ORANGE = "D24E2E", "0E6B45", "8A5B08"
F_GREY, F_GREY2, F_BLUE, F_PINK, F_YEL, F_GRN, F_WARN = "F1F1F2", "F7F7F8", "EFF5FF", "FFF6F5", "FBEFD8", "DFF0E7", "FAE5E3"
TAB_T, TAB_E, TAB_S = "1F6FB2", "C77D1A", "444444"

thin = Side(style="thin", color="D0D2D6")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def font(size=10, bold=False, color=C_TXT):
    return Font(name=FONT, size=size, bold=bold, color=color)


def fill(rgb):
    return PatternFill("solid", fgColor=rgb)


def put(ws, addr, value, *, size=10, bold=False, color=C_TXT, bg=None, fmt=None, h=None, wrap=False, border=True):
    c = ws[addr]
    c.value = value
    c.font = font(size, bold, color)
    if bg:
        c.fill = fill(bg)
    if fmt:
        c.number_format = fmt
    if h or wrap:
        c.alignment = Alignment(horizontal=h, wrap_text=wrap, vertical="center" if wrap else None)
    if border:
        c.border = BORDER
    return c


def header_row(ws, row, labels, start_col=1, height=30):
    for i, lab in enumerate(labels):
        put(ws, f"{L(start_col+i)}{row}", lab, bold=True, color="FFFFFF", bg=C_HEAD, h="center", wrap=True)
    ws.row_dimensions[row].height = height


def q(name):
    return "'" + name + "'"


# ---------- 사업자 정의 ----------
ENT = {
    "T": dict(
        key="T", short="통합", name="따사로운 통합서비스", tab=TAB_T,
        accounts=[  # 결제수단, 구분, 용도, 결제일, 결제계좌, 비고
            ("바우처통장", "통장", "전자바우처 정부지원금 전용", "", "", "산모신생아·가사 등 바우처 정산금이 들어오는 통장"),
            ("경비통장", "통장", "통합서비스 운영경비", "", "", "관리사 용역비·물품·공과금은 여기서 나갑니다"),
            ("삼성카드", "카드", "통합서비스 지출", "매월 __일", "경비통장", ""),
            ("BC카드", "카드", "통합서비스 지출", "매월 __일", "경비통장", ""),
            ("현금", "현금", "현금 지출", "", "", ""),
            (None, None, None, None, None, "← 통합서비스 명의 통장·카드가 더 있으면 여기에 적으십시오"),
            (None, None, None, None, None, ""),
        ],
        voucher=True,
        income=[
            ("바우처수입-산모신생아", "정부지원금 입금분. 바우처통장으로 들어옵니다"),
            ("바우처수입-가사서비스", "정부지원금 입금분"),
            ("바우처수입-영유아", "정부지원금 입금분"),
            ("바우처수입-정리수납", "정부지원금 입금분"),
            ("본인부담금", "이용자가 직접 낸 금액"),
            ("보조금·지원금", "사업 공모·위탁 등 별도 지원금"),
            ("타사업자 정산수입", "평생교육원이 나눠 내는 임차료·공과금 분담분 등. 증빙(계산서·이체확인증)을 반드시 남기십시오"),
            ("이자수입", "통장 이자"),
            ("기타수입", "위에 없는 수입. 적요에 내용을 반드시 적으십시오"),
        ],
        fixed=[
            ("급여(직원)", "내부 직원 급여", ""),
            ("용역비(관리사)", "3.3% 사업소득 지급분", "○"),
            ("4대보험 사업주부담", "", "○"),
            ("퇴직급여", "", ""),
            ("임차료", "사무실 (교육원과 나눠 내면 「타사업자 정산」으로 받으십시오)", ""),
            ("관리비", "", ""),
            ("통신비", "인터넷·전화·휴대폰", ""),
            ("보험료", "배상책임·상해보험", ""),
            ("세무기장료", "", ""),
            ("소프트웨어·구독료", "전산·클라우드·구독 서비스", ""),
            ("차량유지비", "리스·보험·정기분", ""),
        ],
        variable=[
            ("물품구입비", "수건·기저귀·세제 등 서비스 물품"),
            ("소모품비", "사무용품 등"),
            ("교육훈련비", "관리사 보수교육 강사료·교재"),
            ("여비교통비", "출장·주차·유류"),
            ("복리후생비", "직원 식대·경조사"),
            ("접대비", ""),
            ("광고선전비", "홍보물·온라인 광고"),
            ("지급수수료", "은행·카드·플랫폼 수수료"),
            ("수도광열비", "전기·수도·가스"),
            ("세금과공과", "각종 세금·공과금"),
            ("수선비", ""),
            ("타사업자 정산지출", "평생교육원에 보내는 분담금·정산금. 증빙을 반드시 남기십시오"),
            ("잡비", "위에 없는 지출. 적요를 반드시 적으십시오"),
        ],
        examples=[  # 날짜, 결제수단, 계정과목, 거래처, 적요, 수입, 지출, 증빙, 비고
            (dt.date(2026, 9, 1), "바우처통장", "바우처수입-산모신생아", "사회보장정보원", "8월분 바우처 정산 입금", 4820000, None, "이체확인증", ""),
            (dt.date(2026, 9, 1), "경비통장", "계좌대체", "자행이체", "바우처통장 → 경비통장 운영비 이체", 3000000, None, "이체확인증", ""),
            (dt.date(2026, 9, 1), "바우처통장", "계좌대체", "자행이체", "바우처통장 → 경비통장 운영비 이체", None, 3000000, "이체확인증", ""),
            (dt.date(2026, 9, 2), "경비통장", "용역비(관리사)", "관리사 8명", "8월분 용역비 (3.3% 원천징수 후)", None, 2610000, "이체확인증", ""),
            (dt.date(2026, 9, 4), "삼성카드", "임차료", "○○빌딩", "9월 사무실 임차료", None, 850000, "세금계산서", ""),
            (dt.date(2026, 9, 5), "BC카드", "물품구입비", "○○유통", "수건·세제·기저귀 보충", None, 318000, "카드전표", ""),
            (dt.date(2026, 9, 5), "경비통장", "4대보험 사업주부담", "건강보험공단", "8월분 사업주부담금", None, 412000, "이체확인증", ""),
            (dt.date(2026, 9, 10), "경비통장", "통신비", "KT", "9월 인터넷·전화", None, 88000, "세금계산서", ""),
            (dt.date(2026, 9, 15), "경비통장", "타사업자 정산수입", "따사로운 평생교육원", "9월 교육장 임차료 분담분 입금", 300000, None, "이체확인증", "교육원 시트에는 「타사업자 정산지출」로 적혀 있습니다"),
        ],
        stat3=("바우처통장 지출", '=SUMIF($C$15:$C$2014,"바우처통장",$J$15:$J$2014)'),
    ),
    "E": dict(
        key="E", short="교육원", name="따사로운 평생교육원", tab=TAB_E,
        accounts=[
            ("교육원통장", "통장", "수강료·훈련비 수입과 교육원 경비 전부", "", "", "교육원은 이 통장 하나로 수입·지출을 관리합니다"),
            ("현금", "현금", "현금 지출", "", "", ""),
            (None, None, None, None, None, "← 교육원 명의 카드가 있으면 여기에 적으십시오 (구분: 카드, 결제계좌: 교육원통장)"),
            (None, None, None, None, None, ""),
            (None, None, None, None, None, ""),
            (None, None, None, None, None, ""),
            (None, None, None, None, None, ""),
        ],
        voucher=False,
        income=[
            ("수강료수입", "교육원 과정 수강료 (이용자 납부분)"),
            ("훈련비수입(HRD)", "내일배움카드 등 고용노동부 훈련비 지원금"),
            ("교재비수입", "교재 판매분"),
            ("보조금·지원금", "사업 공모·위탁 등 별도 지원금"),
            ("타사업자 정산수입", "통합서비스가 나눠 내는 분담분 등. 증빙을 반드시 남기십시오"),
            ("이자수입", "통장 이자"),
            ("기타수입", "위에 없는 수입. 적요에 내용을 반드시 적으십시오"),
        ],
        fixed=[
            ("급여(직원)", "교육원 직원 급여", ""),
            ("4대보험 사업주부담", "", ""),
            ("퇴직급여", "", ""),
            ("임차료", "교육장 (통합서비스와 나눠 내면 「타사업자 정산지출」로 보내십시오)", ""),
            ("관리비", "", ""),
            ("통신비", "인터넷·전화", ""),
            ("보험료", "", ""),
            ("세무기장료", "", ""),
            ("소프트웨어·구독료", "LMS·전산·구독 서비스", ""),
        ],
        variable=[
            ("강사료", "외부 강사"),
            ("교육·교재비", "교재 구입·인쇄"),
            ("실습재료비", "실습용 재료·소모품"),
            ("소모품비", "사무용품 등"),
            ("여비교통비", "출장·주차·유류"),
            ("복리후생비", "직원 식대·경조사"),
            ("접대비", ""),
            ("광고선전비", "모집 홍보물·온라인 광고"),
            ("지급수수료", "은행·플랫폼 수수료"),
            ("수도광열비", "전기·수도·가스"),
            ("세금과공과", "각종 세금·공과금"),
            ("수선비", ""),
            ("타사업자 정산지출", "통합서비스에 보내는 임차료 분담금 등. 증빙을 반드시 남기십시오"),
            ("잡비", "위에 없는 지출. 적요를 반드시 적으십시오"),
        ],
        examples=[
            (dt.date(2026, 9, 3), "교육원통장", "수강료수입", "김○○ 외 6명", "2026-2기 가사관리사 수강료", 1260000, None, "현금영수증", ""),
            (dt.date(2026, 9, 8), "교육원통장", "강사료", "박○○ 강사", "이유식 실습 4시간", None, 320000, "이체확인증", ""),
            (dt.date(2026, 9, 11), "교육원통장", "교육·교재비", "○○인쇄", "2026-2기 교재 30부", None, 240000, "세금계산서", ""),
            (dt.date(2026, 9, 12), "교육원통장", "계좌대체", "ATM 인출", "교육원통장 → 현금 소액 인출", None, 50000, "이체확인증", ""),
            (dt.date(2026, 9, 12), "현금", "계좌대체", "ATM 인출", "교육원통장 → 현금 소액 인출", 50000, None, "이체확인증", ""),
            (dt.date(2026, 9, 12), "현금", "잡비", "편의점", "교육장 비품 (영수증 분실)", None, 12400, "없음", "증빙 재발급 요청함"),
            (dt.date(2026, 9, 15), "교육원통장", "타사업자 정산지출", "따사로운 통합서비스", "9월 교육장 임차료 분담분 송금", None, 300000, "이체확인증", "통합 시트에는 「타사업자 정산수입」으로 적혀 있습니다"),
        ],
        stat3=("현금 지출", '=SUMIF($C$15:$C$2014,"현금",$J$15:$J$2014)'),
    ),
}
EVIDENCE = ["세금계산서", "계산서", "카드전표", "현금영수증", "이체확인증", "간이영수증", "없음"]

for e in ENT.values():
    s = e["short"]
    e["IN"], e["MO"], e["AC"], e["SET"] = f"{s} ① 거래입력", f"{s} ② 월별결산", f"{s} ③ 통장·카드", f"{s} ④ 설정"

wb = Workbook()
wb.remove(wb.active)

# =====================================================================
# 요약 시트 (맨 앞)
# =====================================================================
S = wb.create_sheet("요약")
S.sheet_properties.tabColor = TAB_S

# =====================================================================
# 사업자별 4개 시트
# =====================================================================
def build_entity(e):
    IN, MO, AC, SET = e["IN"], e["MO"], e["AC"], e["SET"]
    qIN, qMO, qAC, qSET = q(IN), q(MO), q(AC), q(SET)
    ws_in = wb.create_sheet(IN)
    ws_mo = wb.create_sheet(MO)
    ws_ac = wb.create_sheet(AC)
    ws_set = wb.create_sheet(SET)
    for w in (ws_in, ws_mo, ws_ac, ws_set):
        w.sheet_properties.tabColor = e["tab"]

    # ---------------- ④ 설정 ----------------
    ws = ws_set
    put(ws, "A1", f"{e['short']} ④ 설정 — 계정과목표 ({e['name']})", size=15, bold=True, border=False)
    put(ws, "A2", f"여기 목록이 「{IN}」의 드롭다운이 됩니다. 과목을 더하거나 이름을 고치면 입력 화면에 바로 반영됩니다. 이 표는 {e['name']} 전용입니다.", size=9.5, color=C_SUB, border=False)
    put(ws, "A3", "⚠ 계정과목 이름을 바꾸면, 이미 입력된 줄의 이름은 자동으로 바뀌지 않습니다. 바꾸기 전에 ① 시트에서 찾아 바꾸기(Ctrl+H)를 먼저 하십시오.", size=9.5, color=C_RED, border=False)
    if e["voucher"]:
        put(ws, "A4", "⚠ 「바우처통장에서 쓸 수 있는가」에 ○ 표시가 없는 과목을 바우처통장으로 결제하면 ① 시트 점검 칸에 경고가 뜹니다. 지침에 맞게 ○를 켜고 끄십시오 — 지금 켜둔 것은 임시값입니다.", size=9.5, color=C_RED, border=False)
    else:
        put(ws, "A4", "※ 교육원은 통장이 하나라 「계좌대체」는 현금 인출·입금에만 씁니다.", size=9.5, color=C_SUB, border=False)
    header_row(ws, 5, ["대분류", "계정과목", "고정/변동", "설명 · 이런 것이 여기 들어갑니다", "바우처통장에서\n쓸 수 있는가" if e["voucher"] else "(사용 안 함)", "", "증빙"], height=32)
    ws["F5"].fill = PatternFill(fill_type=None); ws["F5"].border = Border()
    r = 6
    rows = []
    for name, desc in e["income"]:
        rows.append(("수입", name, "", desc, ""))
    for name, desc, v in e["fixed"]:
        rows.append(("고정지출", name, "고정", desc, v if e["voucher"] else ""))
    for name, desc in e["variable"]:
        rows.append(("변동지출", name, "변동", desc, ""))
    rows.append(("대체", "계좌대체", "", "통장 사이 이체(통합) / 현금 인출·입금(교육원). 손익에는 잡히지 않고 잔액만 움직입니다", "○" if e["voucher"] else ""))
    assert len(rows) <= 35, len(rows)
    e["n_income"], e["n_fixed"], e["n_var"] = len(e["income"]), len(e["fixed"]), len(e["variable"])
    for i, (cat, name, fv, desc, vch) in enumerate(rows):
        rr = 6 + i
        bg = {"수입": F_BLUE, "고정지출": F_GREY, "변동지출": F_GREY2, "대체": F_YEL}[cat]
        put(ws, f"A{rr}", cat, color=C_SUB, bg=bg, h="center")
        put(ws, f"B{rr}", name, bg=bg)
        put(ws, f"C{rr}", fv, color=C_SUB, bg=bg, h="center")
        put(ws, f"D{rr}", desc, size=9.5, color=C_SUB, bg=bg)
        put(ws, f"E{rr}", vch, color=C_RED, bold=True, bg=bg, h="center")
    for rr in range(6 + len(rows), 41):   # 빈 줄도 테두리 (사용자가 추가 가능)
        for col in "ABCDE":
            put(ws, f"{col}{rr}", None, bg=None)
    for i, ev in enumerate(EVIDENCE):
        put(ws, f"G{6+i}", ev, color=C_SUB, bg=F_GREY2)
    put(ws, "G14", "증빙이 없으면 '없음'을 고르고 비고에 사유를 적으십시오", size=9, color=C_MUTE, border=False)
    for col, w in zip("ABCDEFG", (11, 24, 10, 60, 15, 3, 15)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A6"

    # ---------------- ③ 통장·카드 ----------------
    ws = ws_ac
    put(ws, "A1", f"{e['short']} ③ 통장·카드 ({e['name']} 명의)", size=15, bold=True, border=False)
    put(ws, "A2", f"기초잔액만 넣으시면 나머지는 「{IN}」에서 자동으로 따라옵니다. 여기 A열이 ①의 결제수단 드롭다운 원본입니다. 다른 사업자 통장은 여기 적지 마십시오.", size=9.5, color=C_SUB, border=False)
    put(ws, "D3", "↓ 시작 잔액을 넣으십시오", size=9, color=C_RED, bold=True, border=False)
    header_row(ws, 4, ["결제수단", "구분", "용도", "기초잔액", "수입 합계", "지출 합계", "현재 잔액 / 카드 사용액", "카드 결제일", "카드 결제계좌", "비고"])
    for i, (nm, kind, use, pday, pacct, memo) in enumerate(e["accounts"]):
        rr = 5 + i
        put(ws, f"A{rr}", nm, bold=True)
        put(ws, f"B{rr}", kind, color=C_SUB, h="center")
        put(ws, f"C{rr}", use, size=9.5, color=C_SUB)
        put(ws, f"D{rr}", 0 if nm else None, bg=F_YEL, fmt=NUM, h="right")
        put(ws, f"E{rr}", f'=IF($A{rr}="","",SUMIF({qIN}!$C$15:$C$2014,$A{rr},{qIN}!$I$15:$I$2014))', color=C_GREEN, fmt=NUM, h="right")
        put(ws, f"F{rr}", f'=IF($A{rr}="","",SUMIF({qIN}!$C$15:$C$2014,$A{rr},{qIN}!$J$15:$J$2014))', color="A82219", fmt=NUM, h="right")
        put(ws, f"G{rr}", f'=IF($A{rr}="","",IF($B{rr}="카드",$F{rr},N($D{rr})+$E{rr}-$F{rr}))', bold=True, bg=F_GREY, fmt=NUM, h="right")
        put(ws, f"H{rr}", pday, size=9.5, color=C_SUB, h="center")
        put(ws, f"I{rr}", pacct, size=9.5, color=C_SUB)
        put(ws, f"J{rr}", memo, size=9, color=C_MUTE)
    put(ws, "A12", "통장 합계", bold=True, bg=F_GREY)
    for col in "BC":
        put(ws, f"{col}12", None, bg=F_GREY)
    for col in "DEFG":
        put(ws, f"{col}12", f'=SUMIF($B$5:$B$11,"통장",{col}$5:{col}$11)', size=11, bold=True, bg=F_GREY, fmt=NUM, h="right")
    put(ws, "A13", "카드 사용액 합계", bold=True, bg=F_GREY)
    for col in "BCDE":
        put(ws, f"{col}13", None, bg=F_GREY)
    put(ws, "F13", '=SUMIF($B$5:$B$11,"카드",F$5:F$11)', size=11, bold=True, bg=F_GREY, fmt=NUM, h="right")
    put(ws, "G13", "=$F13", size=11, bold=True, bg=F_GREY, fmt=NUM, h="right")
    notes = [
        "· 잔액이 빨갛게 뜨면 기초잔액을 아직 안 넣었거나, 입력에서 빠진 줄이 있다는 뜻입니다.",
        "· 기초잔액 = 이 회계부를 시작하는 날의 통장 잔액입니다. 한 번만 넣으십시오.",
        "· 카드는 잔액 개념이 없어 「사용액」이 뜹니다. 결제일에 통장에서 빠져나가는 것은 따로 적지 않습니다 — 카드로 쓴 그 줄이 이미 지출로 잡혀 있습니다.",
        "· 통장 실제 잔액과 여기 숫자가 다르면 ① 거래입력에 빠진 줄이 있다는 뜻입니다. 매달 한 번 맞춰 보십시오 — 이것이 회계부의 핵심입니다.",
        f"· 「{e['name']}」 명의가 아닌 통장·카드는 이 표에 넣지 마십시오. 사업자가 다르면 시트가 다릅니다.",
    ]
    for i, t in enumerate(notes):
        put(ws, f"A{15+i}", t, size=9.5, color=C_SUB, border=False)
    ws.conditional_formatting.add("G5:G13", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="A82219", bold=True)))
    for col, w in zip("ABCDEFGHIJ", (16, 8, 34, 14, 14, 14, 20, 12, 14, 52)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"

    # ---------------- ① 거래입력 ----------------
    ws = ws_in
    put(ws, "A1", f"{e['name']} 회계부", size=18, bold=True, border=False)
    ws.row_dimensions[1].height = 26
    put(ws, "G1", f"{e['short']} ① 거래입력 — {e['name']} 통장·카드로 오간 돈만 이 시트에 적습니다. ②③은 자동으로 계산됩니다.", bold=True, color=C_RED, border=False)
    put(ws, "A3", "이번 달", size=11, bold=True, bg=F_GREY)
    put(ws, "B3", '=TEXT(TODAY(),"yyyy년 m월")', size=11, bold=True, color=C_RED, bg=F_GREY, h="left")
    put(ws, "D3", "점검", size=11, bold=True, bg=F_GREY)
    put(ws, "E3", None, bg=F_GREY)
    put(ws, "G3", "쓰는 법", size=11, bold=True, border=False)
    ym = "YEAR(TODAY())*100+MONTH(TODAY())"
    stats = [
        ("수입", f"=SUMIFS($I$15:$I$2014,$O$15:$O$2014,{ym},$P$15:$P$2014,1)"),
        ("지출", f"=SUMIFS($J$15:$J$2014,$O$15:$O$2014,{ym},$P$15:$P$2014,1)"),
        ("차액", "=B4-B5"),
        ("누적 수입", "=SUMIF($P$15:$P$2014,1,$I$15:$I$2014)"),
        ("누적 지출", "=SUMIF($P$15:$P$2014,1,$J$15:$J$2014)"),
        ("입력 줄수", "=COUNT($B$15:$B$2014)"),
    ]
    for i, (lab, f) in enumerate(stats):
        rr = 4 + i
        put(ws, f"A{rr}", lab, color=C_SUB, bg=F_GREY, h="left")
        put(ws, f"B{rr}", f, size=11, bold=True, color=C_GREEN if i != 1 else "A82219", bg=F_GREY, fmt=NUM if i != 5 else "0", h="right")
    checks = [
        ("손볼 줄", '=COUNTIF($L$15:$L$2014,"?*")'),
        ("증빙 없음", '=COUNTIF($K$15:$K$2014,"없음")'),
        e["stat3"],
    ]
    for i, (lab, f) in enumerate(checks):
        rr = 4 + i
        put(ws, f"D{rr}", lab, color=C_SUB, bg=F_GREY, h="left")
        put(ws, f"E{rr}", f, size=11, bold=True, color=C_RED, bg=F_GREY, fmt=NUM, h="right")
    put(ws, "A10", "※ 수입·지출·누적은 「계좌대체」를 뺀 손익 기준입니다.", size=9, color=C_MUTE, border=False)
    howto = [
        f"1. {e['name']} 명의 통장·카드 내역을 내려받아 날짜·거래처·금액을 아래 표에 붙여 넣습니다. 한 줄이 한 거래입니다.",
        "2. 결제수단·계정과목·증빙은 칸을 누르면 목록이 뜹니다. 목록에서 고르십시오 (직접 타이핑하면 ② 결산에서 빠집니다).",
        "3. 대분류·고정/변동은 계정과목을 고르면 저절로 채워집니다. 손대지 마십시오.",
        "4. 수입이면 수입 칸에, 지출이면 지출 칸에만 금액을 적습니다. 둘 다 적지 않습니다.",
        ("5. 통장 사이 이체는 계정과목을 「계좌대체」로 하고 두 줄로 적습니다 — 나가는 통장에 지출, 들어오는 통장에 수입. 손익에는 잡히지 않습니다."
         if e["voucher"] else "5. 현금 인출은 계정과목 「계좌대체」로 두 줄 — 교육원통장에 지출, 현금에 수입. 손익에는 잡히지 않습니다."),
        "6. 다른 사업자와 돈이 오가면(임차료 분담 등) 「계좌대체」가 아니라 「타사업자 정산수입/정산지출」로 적고, 상대 사업자 시트에도 같은 금액을 반대로 적습니다.",
        "7. 「점검」 칸에 글자가 뜨면 그 줄에 빠진 것이 있다는 뜻입니다. 비어 있어야 정상입니다.",
    ]
    for i, t in enumerate(howto):
        put(ws, f"G{4+i}", t, size=9.5, color=C_SUB, border=False)
    put(ws, "G12", f"※ 15행부터 {len(e['examples'])}줄은 형태를 보여드리는 예시입니다. 실제 내역을 넣기 전에 지우십시오.", size=9.5, bold=True, color=C_RED, border=False)

    header_row(ws, 14, ["번호", "날짜", "결제수단", "계정과목", "대분류", "고정/변동", "거래처", "적요 (무엇에 썼는지)", "수입", "지출", "증빙", "점검", "비고"])
    put(ws, "O14", "년월", size=9, color=C_MUTE, border=False)
    put(ws, "P14", "손익", size=9, color=C_MUTE, border=False)

    setB = f"{qSET}!$B$6:$B$40"
    voucher_clause = (f'IF(AND($C{{r}}="바우처통장",$J{{r}}>0,IFERROR(INDEX({qSET}!$E$6:$E$40,MATCH($D{{r}},{setB},0)),"")<>"○"),"바우처통장 지출 — 목적 확인",'
                      if e["voucher"] else "")
    close = ")" if e["voucher"] else ""
    check_tpl = ('=IF($B{r}="","",IF(AND(N($I{r})=0,N($J{r})=0),"금액이 없습니다",IF(AND($I{r}>0,$J{r}>0),"수입·지출 둘 다 적혔습니다",'
                 'IF($C{r}="","결제수단을 고르십시오",IF(ISNA(MATCH($C{r},' + qAC + '!$A$5:$A$11,0)),"③ 통장·카드에 없는 결제수단입니다",'
                 'IF($D{r}="","계정과목을 고르십시오",IF($E{r}="※없는 과목","④ 설정에 없는 과목입니다",'
                 'IF(AND($E{r}="수입",$J{r}>0),"수입 과목인데 지출에 적혔습니다",IF(AND(LEFT($E{r},2)<>"수입",$E{r}<>"대체",$I{r}>0),"지출 과목인데 수입에 적혔습니다",'
                 'IF($K{r}="","증빙을 고르십시오",' + voucher_clause + 'IF($H{r}="","적요를 적으십시오",""))))))))))' + close + ')')

    ex = e["examples"]
    for r in range(NROW0, NROW1 + 1):
        row = ex[r - NROW0] if r - NROW0 < len(ex) else None
        put(ws, f"A{r}", f'=IF($B{r}="","",COUNT($B$15:$B{r}))', size=9.5, color=C_MUTE, h="center")
        put(ws, f"B{r}", row[0] if row else None, fmt="yyyy-mm-dd", h="center")
        put(ws, f"C{r}", row[1] if row else None, size=9.5, color=C_SUB, h="center")
        put(ws, f"D{r}", row[2] if row else None)
        put(ws, f"E{r}", f'=IF($D{r}="","",IFERROR(INDEX({qSET}!$A$6:$A$40,MATCH($D{r},{setB},0)),"※없는 과목"))', size=9.5, color=C_MUTE, h="center")
        put(ws, f"F{r}", f'=IF($D{r}="","",IFERROR(INDEX({qSET}!$C$6:$C$40,MATCH($D{r},{setB},0)),""))', size=9.5, color=C_MUTE, h="center")
        put(ws, f"G{r}", row[3] if row else None)
        put(ws, f"H{r}", row[4] if row else None)
        put(ws, f"I{r}", row[5] if row else None, bold=True, color=C_GREEN, bg=F_BLUE, fmt=NUM, h="right")
        put(ws, f"J{r}", row[6] if row else None, bold=True, color="A82219", bg=F_PINK, fmt=NUM, h="right")
        put(ws, f"K{r}", row[7] if row else None, size=9.5, color=C_SUB, h="center")
        put(ws, f"L{r}", check_tpl.format(r=r), size=9, color=C_ORANGE, bg=F_GREY2)
        put(ws, f"M{r}", row[8] if row else None)
        put(ws, f"O{r}", f'=IF($B{r}="","",YEAR($B{r})*100+MONTH($B{r}))', size=9, color=C_MUTE, border=False)
        put(ws, f"P{r}", f'=IF($B{r}="","",IF($E{r}="대체",0,1))', size=9, color=C_MUTE, border=False)

    dv1 = DataValidation(type="list", formula1=f"={qAC}!$A$5:$A$11", allow_blank=True, showErrorMessage=True,
                         errorTitle="결제수단", error=f"{e['name']} 명의 통장·카드만 고를 수 있습니다. 목록은 ③ 통장·카드 시트에서 고칩니다.")
    dv2 = DataValidation(type="list", formula1=f"={qSET}!$B$6:$B$40", allow_blank=True, showErrorMessage=True,
                         errorTitle="계정과목", error="④ 설정에 있는 과목만 고를 수 있습니다.")
    dv3 = DataValidation(type="list", formula1=f"={qSET}!$G$6:$G$12", allow_blank=True)
    for dv, rng in ((dv1, "C15:C2014"), (dv2, "D15:D2014"), (dv3, "K15:K2014")):
        dv.add(rng); ws.add_data_validation(dv)

    ws.conditional_formatting.add("A15:M2014", FormulaRule(formula=['AND($B15<>"",$L15<>"")'], fill=fill(F_YEL)))
    ws.conditional_formatting.add("A15:M2014", FormulaRule(formula=['$E15="대체"'], font=Font(color=C_MUTE)))
    if e["voucher"]:
        ws.conditional_formatting.add("C15:C2014", FormulaRule(formula=['$C15="바우처통장"'], font=Font(color=C_RED, bold=True)))
    ws.conditional_formatting.add("L15:L2014", FormulaRule(formula=['$L15<>""'], fill=fill(F_WARN), font=Font(color="A82219", bold=True)))
    for col, w in zip("ABCDEFGHIJKLMNOP", (6, 11, 14, 22, 10, 9, 18, 34, 13, 13, 12, 26, 30, 3, 8, 5)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C15"

    # ---------------- ② 월별 결산 ----------------
    ws = ws_mo
    put(ws, "A1", f"{e['short']} ② 월별 결산 ({e['name']})", size=15, bold=True, border=False)
    put(ws, "F1", "연도", size=11, bold=True, border=False, h="right")
    put(ws, "G1", 2026, size=13, bold=True, color=C_RED, bg=F_YEL, fmt="0", h="center")
    put(ws, "H1", "← 이 칸의 연도만 바꾸면 그 해 숫자가 나옵니다", size=9.5, color=C_SUB, border=False)
    put(ws, "A2", f"「{IN}」에서 자동으로 계산됩니다. 손으로 고치지 마십시오. 이 표에는 {e['name']}의 돈만 들어 있습니다.", size=9.5, color=C_SUB, border=False)
    header_row(ws, 4, ["구분", "계정과목"] + [f"{m}월" for m in range(1, 13)] + ["연간 합계", "월평균"])

    def block(r0, title, items, amount_col, bg):
        for i, name in enumerate(items):
            rr = r0 + i
            put(ws, f"A{rr}", title if i == 0 else None, bold=True, bg=bg)
            put(ws, f"B{rr}", name, bg=bg)
            for m in range(1, 13):
                col = L(2 + m)
                put(ws, f"{col}{rr}", f"=SUMIFS({qIN}!${amount_col}$15:${amount_col}$2014,{qIN}!$O$15:$O$2014,$G$1*100+{m},{qIN}!$D$15:$D$2014,$B{rr})",
                    color=C_SUB, bg=bg, fmt=NUM, h="right")
            put(ws, f"O{rr}", f"=SUM($C{rr}:$N{rr})", bold=True, color=C_GREEN, bg=bg, fmt=NUM, h="right")
            put(ws, f"P{rr}", f'=IFERROR(AVERAGEIF($C{rr}:$N{rr},"<>0"),0)', color=C_SUB, bg=bg, fmt=NUM, h="right")
        rs = r0 + len(items)
        put(ws, f"A{rs}", None, bg=F_GREY)
        put(ws, f"B{rs}", f"{title} 소계", bold=True, bg=F_GREY)
        for c in range(3, 17):
            col = L(c)
            put(ws, f"{col}{rs}", f"=SUM({col}{r0}:{col}{rs-1})", bold=True, bg=F_GREY, fmt=NUM, h="right")
        return rs

    inc_names = [n for n, _ in e["income"]]
    fix_names = [n for n, _, _ in e["fixed"]]
    var_names = [n for n, _ in e["variable"]]
    r_inc0 = 5
    r_inc_sub = block(r_inc0, "수입", inc_names, "I", F_BLUE)
    r_fix0 = r_inc_sub + 2
    r_fix_sub = block(r_fix0, "고정지출", fix_names, "J", F_GREY2)
    r_var0 = r_fix_sub + 2
    r_var_sub = block(r_var0, "변동지출", var_names, "J", "FFFFFF")
    r_tot = r_var_sub + 2
    labels = [("총수입", lambda c: f"={c}{r_inc_sub}"),
              ("총지출", lambda c: f"={c}{r_fix_sub}+{c}{r_var_sub}"),
              ("영업이익", lambda c: f"={c}{r_inc_sub}-{c}{r_fix_sub}-{c}{r_var_sub}"),
              ("이익률", lambda c: f'=IF({c}{r_tot}=0,"",({c}{r_tot}-{c}{r_tot+1})/{c}{r_tot})')]
    for i, (lab, fn) in enumerate(labels):
        rr = r_tot + i
        put(ws, f"A{rr}", None, bg=F_GRN)
        put(ws, f"B{rr}", lab, size=11, bold=True, bg=F_GRN)
        for c in range(3, 17):
            col = L(c)
            put(ws, f"{col}{rr}", fn(col), bold=True, color=C_GREEN, bg=F_GRN, fmt="0.0%" if lab == "이익률" else NUM, h="right")
    e["mo_rows"] = dict(tot_in=r_tot, tot_out=r_tot + 1, profit=r_tot + 2)
    put(ws, f"A{r_tot+5}", "※ 계정과목 「계좌대체」는 통장 사이 이동이라 손익에 잡히지 않으므로 이 표에 없습니다. 잔액은 ③ 통장·카드에서 확인하십시오.", size=9.5, color=C_SUB, border=False)
    put(ws, f"A{r_tot+6}", "※ 「월평균」은 금액이 있었던 달만 평균 낸 값입니다. 1년에 한 번 나가는 비용이 12로 나뉘어 작아 보이지 않게 하기 위함입니다.", size=9.5, color=C_SUB, border=False)
    put(ws, f"A{r_tot+7}", "※ 「타사업자 정산수입/정산지출」은 다른 사업자와 실제로 오간 돈이라 손익에 들어갑니다. 두 사업자 시트에 같은 금액이 반대로 적혀 있어야 합니다.", size=9.5, color=C_SUB, border=False)
    for col, w in zip("ABCDEFGHIJKLMNOP", (10, 24) + (12,) * 12 + (14, 12)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C5"


for k in ("T", "E"):
    build_entity(ENT[k])

# =====================================================================
# 요약 시트 내용
# =====================================================================
T, E = ENT["T"], ENT["E"]
ws = S
put(ws, "A1", "따사로운 회계부 — 요약", size=18, bold=True, border=False)
ws.row_dimensions[1].height = 26
put(ws, "A2", "사업자가 둘이므로 회계부도 둘입니다. 이 시트는 두 사업자의 숫자를 나란히 보여줄 뿐, 입력은 각 사업자의 ① 거래입력에서만 합니다. 여기 숫자는 손으로 고치지 마십시오.", size=9.5, color=C_SUB, border=False)

# 구조 안내 (오른쪽)
guide = [
    ("이 파일의 구조", True),
    (f"■ 파란 탭 = {T['name']}  (바우처통장 · 경비통장 · 카드)", False),
    (f"■ 주황 탭 = {E['name']}  (교육원통장 하나)", False),
    ("각 사업자마다 ① 거래입력 → ② 월별결산 → ③ 통장·카드 → ④ 설정, 4장씩입니다.", False),
    ("", False),
    ("지켜야 할 규칙 하나", True),
    ("돈이 어느 사업자 통장·카드에서 오갔는지가 기준입니다. 통합서비스 통장이면 통합 시트, 교육원통장이면 교육원 시트. 「공통」은 없습니다.", False),
    ("함께 쓰는 비용(임차료·공과금 등)은 실제로 낸 사업자가 지출로 적고, 상대가 분담분을 보내면 「타사업자 정산수입 / 정산지출」로 양쪽에 적습니다.", False),
    ("사업자 사이 송금은 반드시 증빙(계산서 또는 이체확인증)을 남기십시오. 세무사에게 넘길 때 두 사업자를 따로 뽑을 수 있습니다.", False),
]
for i, (t, b) in enumerate(guide):
    put(ws, f"G{4+i}", t, size=10 if b else 9.5, bold=b, color=C_TXT if b else C_SUB, border=False)

# 블록 1: 이번 달 / 누적
header_row(ws, 4, ["구분", T["name"], E["name"], "두 사업자 합계 (참고)"], height=30)
summary_rows = [
    ("이번 달 수입", "B4"), ("이번 달 지출", "B5"), ("이번 달 차액", "B6"),
    ("누적 수입", "B7"), ("누적 지출", "B8"),
]
for i, (lab, cell) in enumerate(summary_rows):
    rr = 5 + i
    put(ws, f"A{rr}", lab, color=C_SUB, bg=F_GREY)
    put(ws, f"B{rr}", f"={q(T['IN'])}!{cell}", bold=True, color=C_GREEN, fmt=NUM, h="right")
    put(ws, f"C{rr}", f"={q(E['IN'])}!{cell}", bold=True, color=C_GREEN, fmt=NUM, h="right")
    put(ws, f"D{rr}", f"=B{rr}+C{rr}", bold=True, color=C_SUB, bg=F_GREY2, fmt=NUM, h="right")
put(ws, "A10", "누적 차액", color=C_SUB, bg=F_GREY)
for col in "BCD":
    put(ws, f"{col}10", f"={col}8-{col}9", bold=True, color=C_GREEN, bg=F_GRN, fmt=NUM, h="right")
put(ws, "A11", "손볼 줄 (점검)", color=C_SUB, bg=F_GREY)
put(ws, "B11", f"={q(T['IN'])}!E4", bold=True, color=C_RED, fmt="0", h="right")
put(ws, "C11", f"={q(E['IN'])}!E4", bold=True, color=C_RED, fmt="0", h="right")
put(ws, "D11", "=B11+C11", bold=True, color=C_SUB, bg=F_GREY2, fmt="0", h="right")
put(ws, "A12", "증빙 없음", color=C_SUB, bg=F_GREY)
put(ws, "B12", f"={q(T['IN'])}!E5", bold=True, color=C_RED, fmt="0", h="right")
put(ws, "C12", f"={q(E['IN'])}!E5", bold=True, color=C_RED, fmt="0", h="right")
put(ws, "D12", "=B12+C12", bold=True, color=C_SUB, bg=F_GREY2, fmt="0", h="right")
put(ws, "A13", "※ 이번 달 = 오늘 날짜의 달. 수입·지출은 「계좌대체」를 뺀 손익 기준. 「타사업자 정산」은 한쪽 수입·한쪽 지출이라 합계 열에서는 서로 지워집니다.", size=9, color=C_MUTE, border=False)

# 블록 2: 통장 잔액
put(ws, "A15", "통장·카드 잔액", size=12, bold=True, border=False)
header_row(ws, 16, ["사업자", "결제수단", "구분", "현재 잔액 / 카드 사용액"])
rr = 17
for e in (T, E):
    for i in range(7):
        src = 5 + i
        put(ws, f"A{rr}", f'=IF({q(e["AC"])}!$A{src}="","","{e["short"]}")', color=C_SUB, h="center")
        put(ws, f"B{rr}", f'=IF({q(e["AC"])}!$A{src}="","",{q(e["AC"])}!$A{src})', bold=True)
        put(ws, f"C{rr}", f'=IF({q(e["AC"])}!$A{src}="","",{q(e["AC"])}!$B{src})', color=C_SUB, h="center")
        put(ws, f"D{rr}", f'=IF({q(e["AC"])}!$A{src}="","",{q(e["AC"])}!$G{src})', bold=True, fmt=NUM, h="right")
        rr += 1
    put(ws, f"A{rr}", None, bg=F_GREY)
    put(ws, f"B{rr}", f"{e['short']} 통장 합계", bold=True, bg=F_GREY)
    put(ws, f"C{rr}", None, bg=F_GREY)
    put(ws, f"D{rr}", f"={q(e['AC'])}!$G$12", size=11, bold=True, bg=F_GREY, fmt=NUM, h="right")
    rr += 1
ws.conditional_formatting.add("D17:D32", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="A82219", bold=True)))

# 블록 3: 월별 영업이익
r3 = rr + 2
put(ws, f"A{r3}", "월별 손익 (연도는 각 사업자 ② 월별결산의 G1 칸을 따릅니다)", size=12, bold=True, border=False)
header_row(ws, r3 + 1, ["사업자", "항목"] + [f"{m}월" for m in range(1, 13)] + ["연간 합계"])
rr = r3 + 2
for e in (T, E):
    for lab, key, color in (("총수입", "tot_in", C_SUB), ("총지출", "tot_out", C_SUB), ("영업이익", "profit", C_GREEN)):
        src = e["mo_rows"][key]
        bg = F_GRN if key == "profit" else None
        put(ws, f"A{rr}", e["short"] if key == "tot_in" else None, bold=True, bg=bg)
        put(ws, f"B{rr}", lab, bold=(key == "profit"), bg=bg)
        for c in range(3, 16):
            col = L(c)
            put(ws, f"{col}{rr}", f"={q(e['MO'])}!{col}{src}", bold=(key == "profit"), color=color, bg=bg, fmt=NUM, h="right")
        rr += 1
    e["sum_profit_row"] = rr - 1
put(ws, f"A{rr}", "합계", bold=True, bg=F_GREY)
put(ws, f"B{rr}", "영업이익 (참고)", bold=True, bg=F_GREY)
for c in range(3, 16):
    col = L(c)
    put(ws, f"{col}{rr}", f"={col}{T['sum_profit_row']}+{col}{E['sum_profit_row']}", bold=True, bg=F_GREY, fmt=NUM, h="right")
put(ws, f"A{rr+1}", "※ 「합계」는 대표가 전체를 볼 때만 참고하십시오. 세무·결산은 사업자별 ② 월별결산으로 각각 합니다.", size=9, color=C_MUTE, border=False)

for col, w in zip("ABCDEFGHIJKLMNO", (16, 22, 22, 22, 12, 3, 12, 12, 12, 12, 12, 12, 12, 12, 14)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A4"

wb.active = 0
wb.save(OUT)
print("saved", OUT, [s.title for s in wb.worksheets])
